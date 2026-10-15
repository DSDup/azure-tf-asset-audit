#!/usr/bin/env python3
"""
azure_tf_asset_audit.py - compare an Azure subscription with Terraform state and
produce an Excel asset registry.

Steps (all optional to run separately; typical run does all):
  1. Inventory : `az graph query` (falls back to `az resource list`)  -> inventory.json
  2. Terraform : `terraform show -json` in your TF dir (or a state/show JSON) -> tfstate.json
  3. Compare   : match on lower-cased Azure resource ID
  4. Report    : <out>/azure_asset_registry_<date>.xlsx  (+ raw JSON snapshots)

Examples
  python azure_tf_asset_audit.py --subscription <SUB_ID> --tf-dir C:\\path\\to\\terraform
  python azure_tf_asset_audit.py --subscription <SUB_ID> --tf-dir dirA --tf-dir dirB      # several roots
  python azure_tf_asset_audit.py --inventory-json inv.json --tf-json show.json            # offline

Requires: Python 3.9+, `pip install openpyxl`, Azure CLI (logged in: `az login`),
          Terraform (initialised, with backend access) unless --tf-json is used.
Read-only: nothing here modifies Azure or Terraform state.
"""
import argparse
import datetime as dt
import json
import shutil
import subprocess
import sys
from collections import Counter, defaultdict
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("Missing dependency: pip install openpyxl")

# Azure types that are platform-generated and rarely belong in Terraform.
# Shown as "Ignored" instead of "Unmanaged". Edit freely or pass --ignore-type.
DEFAULT_IGNORE_TYPES = {
    "microsoft.network/networkwatchers",
    "microsoft.network/networkwatchers/flowlogs",
    "microsoft.alertsmanagement/smartdetectoralertrules",
    "microsoft.insights/activitylogalerts",  # remove if you manage these in TF
    "microsoft.security/pricings",
    "microsoft.operationsmanagement/solutions",
    "microsoft.compute/virtualmachines/extensions",
    "microsoft.compute/restorepointcollections",
    "microsoft.portal/dashboards",
    "microsoft.resources/deployments",
}


# --------------------------------------------------------------------------- helpers
def run(cmd, cwd=None):
    exe = shutil.which(cmd[0])
    if not exe:
        sys.exit(f"'{cmd[0]}' not found on PATH.")
    cmd = [exe] + cmd[1:]
    p = subprocess.run(cmd, cwd=cwd, capture_output=True, text=True, shell=False)
    if p.returncode != 0:
        sys.exit(f"Command failed: {' '.join(cmd)}\n{p.stderr.strip()}")
    return p.stdout


def norm(rid):
    return (rid or "").strip().rstrip("/").lower()


# --------------------------------------------------------------------------- inventory
def fetch_inventory(sub):
    """Return list of dict(id,name,type,resourceGroup,location,kind,sku,tags,...)."""
    query = (
        "resources | project id, name, type, resourceGroup, location, kind, "
        "skuName=tostring(sku.name), tags, managedBy, createdTime=tostring(properties.timeCreated), "
        "subscriptionId"
    )
    rows, skip = [], 0
    try:
        while True:
            args = ["az", "graph", "query", "-q", query, "--subscriptions", sub,
                    "--first", "1000", "--skip", str(skip), "-o", "json"]
            data = json.loads(run(args))
            batch = data.get("data", [])
            rows += batch
            if len(batch) < 1000:
                break
            skip += 1000
        return rows
    except SystemExit as e:
        print(f"[warn] Resource Graph failed ({e}); falling back to 'az resource list'.\n"
              f"       (Install with: az extension add --name resource-graph)", file=sys.stderr)
    out = json.loads(run(["az", "resource", "list", "--subscription", sub, "-o", "json"]))
    return [{
        "id": r["id"], "name": r.get("name"), "type": r.get("type"),
        "resourceGroup": r.get("resourceGroup"), "location": r.get("location"),
        "kind": r.get("kind"), "skuName": (r.get("sku") or {}).get("name"),
        "tags": r.get("tags"), "managedBy": r.get("managedBy"),
        "createdTime": r.get("createdTime"), "subscriptionId": sub,
    } for r in out]


# --------------------------------------------------------------------------- terraform
def tf_show(tf_dir):
    out = run(["terraform", "show", "-json"], cwd=tf_dir)
    return json.loads(out)


def walk_module(mod, source):
    for r in mod.get("resources", []):
        if r.get("mode") != "managed":
            continue
        vals = r.get("values") or {}
        yield {
            "address": r.get("address"), "tf_type": r.get("type"),
            "id": vals.get("id") or "", "name": vals.get("name"),
            "source": source,
        }
    for child in mod.get("child_modules", []):
        yield from walk_module(child, source)


def parse_tf(show_json, source):
    root = (show_json.get("values") or {}).get("root_module") or {}
    return list(walk_module(root, source))


# --------------------------------------------------------------------------- compare
def compare(inventory, tf_resources, ignore_types, sub):
    tf_arm = [t for t in tf_resources if norm(t["id"]).startswith("/subscriptions/")]
    tf_other = [t for t in tf_resources if not norm(t["id"]).startswith("/subscriptions/")]
    if sub:
        tf_arm_sub = [t for t in tf_arm if norm(t["id"]).startswith(f"/subscriptions/{sub.lower()}/")]
        tf_other += [dict(t, note="Different subscription") for t in tf_arm if t not in tf_arm_sub]
        tf_arm = tf_arm_sub

    tf_by_id = defaultdict(list)
    for t in tf_arm:
        tf_by_id[norm(t["id"])].append(t)
    tf_ids = list(tf_by_id)

    az_ids = {norm(r["id"]) for r in inventory}
    registry, matched_tf = [], set()

    for r in inventory:
        rid = norm(r["id"])
        rtype = (r.get("type") or "").lower()
        hits = tf_by_id.get(rid)
        via = "exact"
        if not hits:
            # Terraform may manage a child object (e.g. subnet, diagnostic setting) of this resource
            child = [i for i in tf_ids if i.startswith(rid + "/")]
            if child:
                hits = [t for i in child for t in tf_by_id[i]]
                via = "child resource"
        if hits:
            status = "Managed"
            matched_tf.update(norm(t["id"]) for t in hits)
            addr = "; ".join(sorted({h["address"] for h in hits}))[:500]
        elif rtype in ignore_types:
            status, addr, via = "Ignored", "", "type in ignore list"
        else:
            status, addr, via = "Unmanaged", "", ""
        registry.append({**r, "status": status, "tf_address": addr, "match": via})

    # TF resources not matching anything in Azure inventory
    missing = []
    for rid, items in tf_by_id.items():
        if rid in matched_tf or rid in az_ids:
            continue
        # skip children of inventory resources (they were matched above via prefix)
        if any(rid.startswith(a + "/") for a in az_ids):
            continue
        for t in items:
            missing.append({**t, "note": "In Terraform state but not found in Azure (deleted outside TF, "
                                         "non-ARG resource type, or stale state)"})
    return registry, missing, tf_other


# --------------------------------------------------------------------------- excel
HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(bold=True, color="FFFFFF")
STATUS_FILL = {
    "Managed": PatternFill("solid", fgColor="E2EFDA"),
    "Unmanaged": PatternFill("solid", fgColor="FCE4D6"),
    "Ignored": PatternFill("solid", fgColor="EDEDED"),
}


def write_sheet(wb, title, headers, rows, status_col=None):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in ws[1]:
        c.fill, c.font = HDR_FILL, HDR_FONT
        c.alignment = Alignment(vertical="center")
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions
    for i, h in enumerate(headers, 1):
        width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows[:500] if i - 1 < len(r)])
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 60)
    if status_col is not None:
        for r in ws.iter_rows(min_row=2):
            f = STATUS_FILL.get(r[status_col].value)
            if f:
                r[status_col].fill = f
    return ws


def tags_str(t):
    return "; ".join(f"{k}={v}" for k, v in sorted((t or {}).items()))


def build_workbook(path, sub, registry, missing, tf_other, tf_total, tf_sources):
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    counts = Counter(r["status"] for r in registry)
    total = len(registry)
    cov = counts["Managed"] / max(total - counts["Ignored"], 1)
    lines = [
        ("Azure vs Terraform asset registry", ""),
        ("Generated (UTC)", dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%d %H:%M")),
        ("Subscription", sub or "(from inventory file)"),
        ("Terraform sources", ", ".join(tf_sources) or "(none)"),
        ("", ""),
        ("Azure resources (inventory)", total),
        ("  Managed by Terraform", counts["Managed"]),
        ("  Unmanaged (in Azure, not in TF)", counts["Unmanaged"]),
        ("  Ignored (platform / ignore list)", counts["Ignored"]),
        ("Terraform coverage (excl. ignored)", f"{cov:.1%}"),
        ("", ""),
        ("Terraform managed resources in state", tf_total),
        ("  In state but NOT found in Azure", len(missing)),
        ("  Non-Azure-ARM / other-subscription (azuread, random, ...)", len(tf_other)),
        ("", ""),
        ("Notes", "Match = lower-cased Azure resource ID. TF child objects (subnets, diagnostic settings, "
                  "etc.) count as managing their parent. Role assignments, policy and Entra objects are "
                  "not in Resource Graph 'resources' and are out of scope."),
    ]
    for l in lines:
        ws.append(list(l))
    ws["A1"].font = Font(bold=True, size=14)
    ws.column_dimensions["A"].width = 58
    ws.column_dimensions["B"].width = 80
    for r in ws.iter_rows(min_row=2):
        r[0].font = Font(bold=True)
        r[1].alignment = Alignment(wrap_text=True, horizontal="left", vertical="top")

    reg_rows = [[r["status"], r.get("name"), r.get("type"), r.get("resourceGroup"), r.get("location"),
                 r.get("kind"), r.get("skuName"), tags_str(r.get("tags")), r.get("createdTime"),
                 r.get("tf_address"), r.get("match"), r.get("managedBy"), r["id"]] for r in registry]
    reg_rows.sort(key=lambda x: (str(x[3]).lower(), str(x[2]).lower(), str(x[1]).lower()))
    hdr = ["Status", "Name", "Type", "Resource Group", "Location", "Kind", "SKU", "Tags", "Created",
           "Terraform address", "Match", "Managed by", "Resource ID"]
    write_sheet(wb, "Registry", hdr, reg_rows, status_col=0)
    write_sheet(wb, "Unmanaged", hdr, [r for r in reg_rows if r[0] == "Unmanaged"], status_col=0)

    # by type
    by_type = defaultdict(Counter)
    for r in registry:
        by_type[(r.get("type") or "").lower()][r["status"]] += 1
    rows = [[t, sum(c.values()), c["Managed"], c["Unmanaged"], c["Ignored"]] for t, c in sorted(by_type.items())]
    write_sheet(wb, "By type", ["Type", "Total", "Managed", "Unmanaged", "Ignored"], rows)

    by_rg = defaultdict(Counter)
    for r in registry:
        by_rg[(r.get("resourceGroup") or "").lower()][r["status"]] += 1
    rows = [[g, sum(c.values()), c["Managed"], c["Unmanaged"], c["Ignored"]] for g, c in sorted(by_rg.items())]
    write_sheet(wb, "By resource group", ["Resource group", "Total", "Managed", "Unmanaged", "Ignored"], rows)

    write_sheet(wb, "In TF not in Azure",
                ["Terraform address", "TF type", "Name", "Source", "Resource ID", "Note"],
                [[m["address"], m["tf_type"], m["name"], m["source"], m["id"], m["note"]] for m in missing])
    write_sheet(wb, "TF non-Azure", ["Terraform address", "TF type", "Name", "Source", "ID", "Note"],
                [[m["address"], m["tf_type"], m["name"], m["source"], m["id"], m.get("note", "")] for m in tf_other])
    wb.save(path)


# --------------------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--subscription", "-s", help="Azure subscription ID (default: current az account)")
    ap.add_argument("--tf-dir", action="append", default=[], help="Terraform root dir (repeatable)")
    ap.add_argument("--tf-json", action="append", default=[],
                    help="Pre-exported `terraform show -json` file (repeatable)")
    ap.add_argument("--inventory-json", help="Use an existing inventory JSON instead of calling az")
    ap.add_argument("--out", default="output", help="Output directory (default: ./output)")
    ap.add_argument("--ignore-type", action="append", default=[],
                    help="Extra Azure resource type to ignore, e.g. microsoft.insights/components")
    ap.add_argument("--no-default-ignores", action="store_true", help="Don't use the built-in ignore list")
    a = ap.parse_args()

    out = Path(a.out)
    out.mkdir(parents=True, exist_ok=True)
    stamp = dt.datetime.now().strftime("%Y%m%d_%H%M")

    sub = a.subscription
    if not a.inventory_json and not sub:
        sub = run(["az", "account", "show", "--query", "id", "-o", "tsv"]).strip()
    if a.inventory_json:
        inventory = json.loads(Path(a.inventory_json).read_text(encoding="utf-8"))
    else:
        print(f"Fetching inventory for {sub} ...")
        inventory = fetch_inventory(sub)
        (out / f"inventory_{stamp}.json").write_text(json.dumps(inventory, indent=2), encoding="utf-8")
    print(f"  {len(inventory)} Azure resources")

    tf_resources, sources = [], []
    for d in a.tf_dir:
        print(f"Reading Terraform state in {d} ...")
        js = tf_show(d)
        (out / f"tfshow_{Path(d).name}_{stamp}.json").write_text(json.dumps(js), encoding="utf-8")
        tf_resources += parse_tf(js, d)
        sources.append(d)
    for f in a.tf_json:
        tf_resources += parse_tf(json.loads(Path(f).read_text(encoding="utf-8")), f)
        sources.append(f)
    if not sources:
        print("[warn] No Terraform source given: everything will be 'Unmanaged'.", file=sys.stderr)
    print(f"  {len(tf_resources)} Terraform managed resources")

    ignore = set(a.ignore_type) | (set() if a.no_default_ignores else DEFAULT_IGNORE_TYPES)
    ignore = {i.lower() for i in ignore}
    registry, missing, tf_other = compare(inventory, tf_resources, ignore, sub)

    xlsx = out / f"azure_asset_registry_{stamp}.xlsx"
    build_workbook(xlsx, sub, registry, missing, tf_other, len(tf_resources), sources)
    c = Counter(r["status"] for r in registry)
    print(f"\nManaged {c['Managed']} | Unmanaged {c['Unmanaged']} | Ignored {c['Ignored']} | "
          f"In TF not Azure {len(missing)}")
    print(f"Report: {xlsx}")


if __name__ == "__main__":
    main()
